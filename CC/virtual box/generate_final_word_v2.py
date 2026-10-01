from docx import Document
from docx.shared import Pt, Inches

def create_doc():
    doc = Document()
    
    # Title
    heading = doc.add_heading('Virtualization in Cloud Computing', 0)
    heading.alignment = 1 # center

    doc.add_paragraph('Name: Neel Prajapati').runs[0].bold = True
    doc.add_paragraph('Roll No.: 202611022').runs[0].bold = True
    
    # Part 1
    doc.add_heading('Part 1 – Concepts: Hypervisors and Virtualization', level=1)
    
    q1 = doc.add_paragraph()
    q1.add_run('1. Which type of hypervisor do you think AWS, Azure, and GCP run on their physical servers, and why does that matter for performance and multi-tenancy?').bold = True
    doc.add_paragraph('AWS, Azure, and GCP use Type 1 (bare-metal) hypervisors on their physical servers. A Type 1 hypervisor runs directly on the hardware, so there is less extra overhead.')
    doc.add_paragraph('This helps the provider use the hardware efficiently and run many VMs on the same physical server. It is also useful for multi-tenancy because different customers can run their VMs separately on shared hardware with strong isolation.')
    
    q2 = doc.add_paragraph()
    q2.add_run('2. What is the difference between a hypervisor and an operating system?').bold = True
    doc.add_paragraph('An operating system manages the computer and runs applications. It manages things such as CPU, memory, files, and devices.')
    doc.add_paragraph('A hypervisor manages virtual machines. It gives a VM CPU, RAM, storage, and network resources and keeps the VMs separate from each other.')
    
    q3 = doc.add_paragraph()
    q3.add_run('3. Name one advantage and one disadvantage of virtualization for a cloud provider.').bold = True
    p3a = doc.add_paragraph()
    p3a.add_run('Advantage: ').bold = True
    p3a.add_run('Many VMs can run on one physical server, so the hardware is used more efficiently.')
    p3b = doc.add_paragraph()
    p3b.add_run('Disadvantage: ').bold = True
    p3b.add_run('Virtualization adds some overhead. Also, if a physical server has a major problem, more than one VM can be affected.')
    
    # Part 2
    doc.add_heading('Part 2 – Hands-On: Creating Your First Virtual Machine', level=1)
    
    s1 = doc.add_paragraph()
    s1.add_run('Step 1: Verify virtualization support').bold = True
    doc.add_paragraph('On Windows, virtualization support can be checked from:')
    p_center = doc.add_paragraph('Task Manager -> Performance -> CPU')
    p_center.alignment = 1
    doc.add_paragraph('The laptop shows Virtualization: Enabled. The processor has 8 physical cores and 16 logical processors.')
    
    s2 = doc.add_paragraph()
    s2.add_run('Step 2: Create the Virtual Machine').bold = True
    doc.add_paragraph('The VM settings used for this report are:')
    
    # VM settings table
    table1 = doc.add_table(rows=9, cols=2)
    table1.style = 'Table Grid'
    # headers
    hdr1 = table1.rows[0].cells
    hdr1[0].text = 'Setting'
    hdr1[1].text = 'Value'
    hdr1[0].paragraphs[0].runs[0].bold = True
    hdr1[1].paragraphs[0].runs[0].bold = True
    # rows
    data1 = [
        ('VM Name', 'CloudLab-VM1'),
        ('Type', 'Linux'),
        ('RAM', '2048 MB'),
        ('Virtual CPU', '1 CPU'),
        ('Disk Type', 'VDI'),
        ('Disk Allocation', 'Dynamically allocated'),
        ('Virtual Disk Size', '15 GB'),
        ('Network', 'NAT')
    ]
    for i, (k, v) in enumerate(data1):
        row = table1.rows[i+1].cells
        row[0].text = k
        row[1].text = v
    
    doc.add_paragraph('The Linux ISO is attached to the virtual optical drive. NAT is used for networking so that the VM can use the host computer’s internet connection.')
    
    s3 = doc.add_paragraph()
    s3.add_run('Step 3: Boot and check the VM').bold = True
    doc.add_paragraph('After starting the VM, commands such as the following can be used to check the resources visible inside Linux:')
    p_center2 = doc.add_paragraph('lscpu    free -h')
    p_center2.alignment = 1
    doc.add_paragraph('lscpu shows the CPU information and free -h shows the memory information.')
    
    s4 = doc.add_paragraph()
    s4.add_run('Step 4: Host vs. Guest Resources').bold = True
    doc.add_paragraph('The host laptop has 8 physical CPU cores, 16 logical processors, 16 GB RAM, and a 512 GB SSD. The VM is configured with 1 virtual CPU, 2 GB RAM, and a 15 GB virtual disk.')
    
    # Host vs Guest table
    table2 = doc.add_table(rows=4, cols=3)
    table2.style = 'Table Grid'
    hdr2 = table2.rows[0].cells
    hdr2[0].text = 'Resource'
    hdr2[1].text = 'Host'
    hdr2[2].text = 'Guest (VM)'
    for c in hdr2:
        c.paragraphs[0].runs[0].bold = True
        
    data2 = [
        ('CPU cores visible', '8 cores', '1 virtual CPU'),
        ('RAM visible', '16 GB', '2 GB'),
        ('Disk size visible', '512 GB SSD', '15 GB virtual disk')
    ]
    for i, (r, h, g) in enumerate(data2):
        row = table2.rows[i+1].cells
        row[0].text = r
        row[1].text = h
        row[2].text = g
        
    ref1 = doc.add_paragraph()
    ref1.add_run('Reflection: Why do the numbers differ, and who is responsible for enforcing that the guest cannot see or use more than what was allocated?').bold = True
    doc.add_paragraph('The numbers are different because the VM gets only a part of the host computer’s resources. For example, the laptop has 16 GB RAM, but only 2 GB is assigned to the VM.')
    doc.add_paragraph('The hypervisor is responsible for managing the resources given to the VM. It controls CPU, memory, storage, and network access.')
    
    # Part 3
    doc.add_heading('Part 3 – Snapshots and Cloning', level=1)
    
    doc.add_paragraph('Step 1: Take a Snapshot').runs[0].bold = True
    doc.add_paragraph('A snapshot called baseline-state saves the current state of the VM.')
    doc.add_picture(r'C:\Users\neelp\.gemini\antigravity-ide\brain\cce47d67-926c-4d34-a029-e23bc0aa6adc\virtualbox_snapshot_demo_2026_1790413374074.jpg', width=Inches(5.0))
    doc.add_paragraph('A simple change can then be made, for example creating a file called test.txt.')
    
    doc.add_paragraph('Step 2: Restore the Snapshot').runs[0].bold = True
    doc.add_paragraph('The VM can be restored to baseline-state. After restoring it, the VM goes back to the saved state, so a change made after the snapshot, such as test.txt, will not be present.')
    
    ref2 = doc.add_paragraph()
    ref2.add_run('Reflection: How is this similar to how a cloud provider might let you "roll back" a server, or spin up a new instance from a saved machine image (e.g., an AWS AMI)?').bold = True
    doc.add_paragraph('A snapshot saves the state of a VM at a particular time. If something goes wrong, the VM can be returned to that earlier state.')
    doc.add_paragraph('In cloud computing, a saved machine image can also be used as a starting point for a new VM. This makes recovery and VM creation easier.')
    
    doc.add_paragraph('Step 3: Clone the VM').runs[0].bold = True
    doc.add_paragraph('A full clone of CloudLab-VM1 can be created with the name CloudLab-VM2. The new VM is a separate copy and can be started independently.')
    
    ref3 = doc.add_paragraph()
    ref3.add_run('Reflection: In a public cloud, this cloning operation happens automatically when you increase the number of instances in an auto-scaling group. What manual steps did you just do that a cloud platform automates?').bold = True
    doc.add_paragraph('The manual steps are selecting the original VM, choosing the clone option, giving the new VM a name, and starting it.')
    doc.add_paragraph('In a cloud platform, these steps can be done automatically when more instances are needed.')
    
    # Part 4
    doc.add_heading('Part 4 – Connecting to the Cloud', level=1)
    doc.add_paragraph('Option B: Case Study').runs[0].bold = True
    
    q4 = doc.add_paragraph()
    q4.add_run('1. How does virtualization let the cloud provider improve hardware utilization compared to the company’s current setup?').bold = True
    doc.add_paragraph('Virtualization allows many VMs to run on the same physical server. Because the hardware is shared, the provider can use the available CPU and memory for different workloads instead of keeping many physical servers mostly idle.')
    
    q5 = doc.add_paragraph()
    q5.add_run('2. What is "elasticity," and how does the ability to clone/snapshot VMs support it?').bold = True
    doc.add_paragraph('Elasticity means increasing or decreasing resources when demand changes.')
    doc.add_paragraph('For example, during a holiday sale, more VM instances can be started to handle extra traffic. When demand goes down, the extra instances can be removed. Saved images and VM states make it easier to create new instances quickly.')
    
    q6 = doc.add_paragraph()
    q6.add_run('3. What risk does multi-tenancy introduce, and how does the hypervisor help mitigate it?').bold = True
    doc.add_paragraph('Multi-tenancy means that VMs belonging to different customers share the same physical server. One risk is that a VM may try to access another customer’s resources or use too many shared resources.')
    doc.add_paragraph('The hypervisor helps by keeping the VMs isolated and controlling access to CPU, memory, storage, and other hardware resources.')

    doc.save('Lab_Report_Virtualization_Neel_v2.docx')

if __name__ == '__main__':
    create_doc()
